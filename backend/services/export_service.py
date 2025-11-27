import pandas as pd
import json
from datetime import datetime
from typing import Dict, Optional, Any
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import os

class ExportService:
    def __init__(self):
        self.export_dir = "exports"
        os.makedirs(self.export_dir, exist_ok=True)
    
    async def export_data(
        self, 
        chef_id: int, 
        export_type: str, 
        format_type: str = "excel",
        date_range: Optional[Dict[str, datetime]] = None,
        filters: Optional[Dict] = None,
        db: Session = None
    ) -> str:
        """Export data to Excel or CSV format with professional formatting"""
        
        from models import (
            SupplierDB, SupplierProductDB, PurchaseOrderDB, InventoryItemDB,
            AIConsultantRecommendationDB, InventoryAlertDB
        )
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        if export_type == "suppliers":
            return await self._export_suppliers(chef_id, format_type, timestamp, db)
        elif export_type == "inventory":
            return await self._export_inventory(chef_id, format_type, timestamp, db)
        elif export_type == "orders":
            return await self._export_orders(chef_id, format_type, timestamp, db)
        elif export_type == "recommendations":
            return await self._export_recommendations(chef_id, format_type, timestamp, db)
        elif export_type == "alerts":
            return await self._export_alerts(chef_id, format_type, timestamp, db)
        else:
            raise ValueError(f"Unsupported export type: {export_type}")
    
    async def _export_suppliers(self, chef_id: int, format_type: str, timestamp: str, db: Session) -> str:
        """Export supplier data"""
        from models import SupplierDB, SupplierProductDB
        
        suppliers = db.query(SupplierDB).filter(SupplierDB.chef_id == chef_id).all()
        
        data = []
        for supplier in suppliers:
            # Get product count for each supplier
            product_count = db.query(SupplierProductDB).filter(
                SupplierProductDB.supplier_id == supplier.id
            ).count()
            
            data.append({
                "Supplier Name": supplier.name,
                "Contact Person": supplier.contact_person or "N/A",
                "Email": supplier.email or "N/A",
                "Phone": supplier.phone or "N/A",
                "Country": supplier.country,
                "City": supplier.city or "N/A",
                "Payment Terms": supplier.payment_terms or "N/A",
                "Delivery Time": supplier.delivery_time or "N/A",
                "Minimum Order": supplier.minimum_order,
                "Rating": supplier.rating,
                "Specialties": ", ".join(supplier.specialties) if supplier.specialties else "N/A",
                "Brands Carried": ", ".join(supplier.brands_carried) if supplier.brands_carried else "N/A",
                "Total Orders": supplier.total_orders,
                "Total Spent (TND)": supplier.total_spent,
                "Last Order Date": supplier.last_order_date.strftime("%Y-%m-%d") if supplier.last_order_date else "N/A",
                "Product Count": product_count,
                "Status": "Active" if supplier.is_active else "Inactive",
                "Verified": "Yes" if supplier.is_verified else "No",
                "Created Date": supplier.created_at.strftime("%Y-%m-%d")
            })
        
        df = pd.DataFrame(data)
        
        if format_type == "csv":
            filename = f"suppliers_export_{timestamp}.csv"
            filepath = os.path.join(self.export_dir, filename)
            df.to_csv(filepath, index=False)
        else:  # Excel
            filename = f"suppliers_export_{timestamp}.xlsx"
            filepath = os.path.join(self.export_dir, filename)
            
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Suppliers', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Suppliers']
                
                # Format header
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add summary sheet
                summary_data = {
                    "Metric": [
                        "Total Suppliers",
                        "Active Suppliers", 
                        "Verified Suppliers",
                        "Total Spending (TND)",
                        "Average Rating",
                        "Average Orders per Supplier"
                    ],
                    "Value": [
                        len(suppliers),
                        len([s for s in suppliers if s.is_active]),
                        len([s for s in suppliers if s.is_verified]),
                        sum(s.total_spent for s in suppliers),
                        sum(s.rating for s in suppliers) / len(suppliers) if suppliers else 0,
                        sum(s.total_orders for s in suppliers) / len(suppliers) if suppliers else 0
                    ]
                }
                
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Format summary sheet
                summary_sheet = writer.sheets['Summary']
                for cell in summary_sheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
        
        return filepath
    
    async def _export_inventory(self, chef_id: int, format_type: str, timestamp: str, db: Session) -> str:
        """Export inventory data"""
        from models import InventoryItemDB, InventoryAlertDB
        
        inventory_items = db.query(InventoryItemDB).filter(InventoryItemDB.chef_id == chef_id).all()
        
        data = []
        for item in inventory_items:
            # Calculate stock status
            if item.current_stock <= item.critical_stock:
                stock_status = "CRITICAL"
            elif item.current_stock <= item.minimum_stock:
                stock_status = "LOW"
            else:
                stock_status = "OK"
            
            # Calculate days until stock out
            if item.average_monthly_consumption > 0:
                daily_consumption = item.average_monthly_consumption / 30
                days_until_out = item.current_stock / daily_consumption if daily_consumption > 0 else 999
            else:
                days_until_out = 999
            
            data.append({
                "Item Name": item.item_name,
                "Category": item.category or "N/A",
                "Brand": item.brand or "N/A",
                "Current Stock": item.current_stock,
                "Unit": item.unit,
                "Unit Cost (TND)": item.unit_cost,
                "Total Value (TND)": item.current_stock * item.unit_cost,
                "Stock Status": stock_status,
                "Days Until Stock Out": int(days_until_out),
                "Minimum Stock": item.minimum_stock,
                "Critical Stock": item.critical_stock,
                "Monthly Consumption": item.average_monthly_consumption,
                "Storage Location": item.storage_location or "N/A",
                "Shelf Life (Days)": item.shelf_life_days or "N/A",
                "Expiry Date": item.expiry_date.strftime("%Y-%m-%d") if item.expiry_date else "N/A",
                "Last Updated": item.last_updated.strftime("%Y-%m-%d %H:%M"),
                "Status": item.status
            })
        
        df = pd.DataFrame(data)
        
        if format_type == "csv":
            filename = f"inventory_export_{timestamp}.csv"
            filepath = os.path.join(self.export_dir, filename)
            df.to_csv(filepath, index=False)
        else:  # Excel
            filename = f"inventory_export_{timestamp}.xlsx"
            filepath = os.path.join(self.export_dir, filename)
            
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Inventory', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Inventory']
                
                # Format header
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add summary sheet
                low_stock_items = len([item for item in inventory_items if item.current_stock <= item.minimum_stock])
                critical_items = len([item for item in inventory_items if item.current_stock <= item.critical_stock])
                total_value = sum(item.current_stock * item.unit_cost for item in inventory_items)
                
                summary_data = {
                    "Metric": [
                        "Total Items",
                        "Low Stock Items",
                        "Critical Items",
                        "Total Inventory Value (TND)",
                        "Average Unit Cost (TND)",
                        "Items with Expiry Dates"
                    ],
                    "Value": [
                        len(inventory_items),
                        low_stock_items,
                        critical_items,
                        total_value,
                        sum(item.unit_cost for item in inventory_items) / len(inventory_items) if inventory_items else 0,
                        len([item for item in inventory_items if item.expiry_date])
                    ]
                }
                
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Format summary sheet
                summary_sheet = writer.sheets['Summary']
                for cell in summary_sheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
        
        return filepath
    
    async def _export_orders(self, chef_id: int, format_type: str, timestamp: str, db: Session) -> str:
        """Export purchase orders data"""
        from models import PurchaseOrderDB, PurchaseOrderItemDB, SupplierDB, SupplierProductDB
        
        orders = db.query(PurchaseOrderDB).filter(PurchaseOrderDB.chef_id == chef_id).all()
        
        data = []
        for order in orders:
            supplier = db.query(SupplierDB).filter(SupplierDB.id == order.supplier_id).first()
            order_items = db.query(PurchaseOrderItemDB).filter(PurchaseOrderItemDB.order_id == order.id).all()
            
            for item in order_items:
                product = db.query(SupplierProductDB).filter(SupplierProductDB.id == item.product_id).first()
                
                data.append({
                    "Order Number": order.order_number,
                    "Order Date": order.order_date.strftime("%Y-%m-%d %H:%M"),
                    "Supplier Name": supplier.name if supplier else "N/A",
                    "Product Name": product.product_name if product else "N/A",
                    "Brand": product.brand if product else "N/A",
                    "Quantity": item.quantity,
                    "Unit Price (TND)": item.unit_price,
                    "Total Price (TND)": item.total_price,
                    "Received Quantity": item.received_quantity,
                    "Order Status": order.status,
                    "Priority": order.priority,
                    "Payment Status": order.payment_status,
                    "Requested Delivery": order.requested_delivery_date.strftime("%Y-%m-%d") if order.requested_delivery_date else "N/A",
                    "Actual Delivery": order.actual_delivery_date.strftime("%Y-%m-%d") if order.actual_delivery_date else "N/A",
                    "Total Order Amount (TND)": order.total_amount,
                    "Currency": order.currency,
                    "Notes": item.notes or "N/A"
                })
        
        df = pd.DataFrame(data)
        
        if format_type == "csv":
            filename = f"orders_export_{timestamp}.csv"
            filepath = os.path.join(self.export_dir, filename)
            df.to_csv(filepath, index=False)
        else:  # Excel
            filename = f"orders_export_{timestamp}.xlsx"
            filepath = os.path.join(self.export_dir, filename)
            
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Orders', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Orders']
                
                # Format header
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add summary sheet
                pending_orders = len([o for o in orders if o.status == "pending"])
                delivered_orders = len([o for o in orders if o.status == "delivered"])
                total_order_value = sum(o.total_amount for o in orders)
                
                summary_data = {
                    "Metric": [
                        "Total Orders",
                        "Pending Orders",
                        "Delivered Orders",
                        "Total Order Value (TND)",
                        "Average Order Value (TND)",
                        "Orders This Month"
                    ],
                    "Value": [
                        len(orders),
                        pending_orders,
                        delivered_orders,
                        total_order_value,
                        total_order_value / len(orders) if orders else 0,
                        len([o for o in orders if o.order_date.month == datetime.now().month])
                    ]
                }
                
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Format summary sheet
                summary_sheet = writer.sheets['Summary']
                for cell in summary_sheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
        
        return filepath
    
    async def _export_recommendations(self, chef_id: int, format_type: str, timestamp: str, db: Session) -> str:
        """Export AI consultant recommendations"""
        from models import AIConsultantRecommendationDB
        
        recommendations = db.query(AIConsultantRecommendationDB).filter(
            AIConsultantRecommendationDB.chef_id == chef_id
        ).all()
        
        data = []
        for rec in recommendations:
            data.append({
                "Title": rec.title,
                "Type": rec.recommendation_type,
                "Description": rec.description,
                "Cost Savings Potential (TND)": rec.cost_savings_potential,
                "Quality Score": rec.quality_score,
                "Reliability Score": rec.reliability_score,
                "Value Score": rec.value_score,
                "Implementation Priority": rec.implementation_priority,
                "Estimated Time": rec.estimated_implementation_time or "N/A",
                "Implemented": "Yes" if rec.is_implemented else "No",
                "Created Date": rec.created_at.strftime("%Y-%m-%d %H:%M"),
                "Expires": rec.expires_at.strftime("%Y-%m-%d") if rec.expires_at else "N/A",
                "Implementation Notes": rec.implementation_notes or "N/A"
            })
        
        df = pd.DataFrame(data)
        
        if format_type == "csv":
            filename = f"recommendations_export_{timestamp}.csv"
            filepath = os.path.join(self.export_dir, filename)
            df.to_csv(filepath, index=False)
        else:  # Excel
            filename = f"recommendations_export_{timestamp}.xlsx"
            filepath = os.path.join(self.export_dir, filename)
            
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Recommendations', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Recommendations']
                
                # Format header
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="9B59B6", end_color="9B59B6", fill_type="solid")
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add summary sheet
                implemented_count = len([r for r in recommendations if r.is_implemented])
                total_savings = sum(r.cost_savings_potential for r in recommendations)
                
                summary_data = {
                    "Metric": [
                        "Total Recommendations",
                        "Implemented Recommendations",
                        "Pending Recommendations",
                        "Total Potential Savings (TND)",
                        "Average Quality Score",
                        "Average Reliability Score"
                    ],
                    "Value": [
                        len(recommendations),
                        implemented_count,
                        len(recommendations) - implemented_count,
                        total_savings,
                        sum(r.quality_score for r in recommendations) / len(recommendations) if recommendations else 0,
                        sum(r.reliability_score for r in recommendations) / len(recommendations) if recommendations else 0
                    ]
                }
                
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Format summary sheet
                summary_sheet = writer.sheets['Summary']
                for cell in summary_sheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
        
        return filepath
    
    async def _export_alerts(self, chef_id: int, format_type: str, timestamp: str, db: Session) -> str:
        """Export inventory alerts"""
        from models import InventoryAlertDB, InventoryItemDB
        
        alerts = db.query(InventoryAlertDB).filter(InventoryAlertDB.chef_id == chef_id).all()
        
        data = []
        for alert in alerts:
            item = db.query(InventoryItemDB).filter(InventoryItemDB.id == alert.inventory_item_id).first()
            
            data.append({
                "Alert Type": alert.alert_type,
                "Item Name": item.item_name if item else "N/A",
                "Category": item.category if item else "N/A",
                "Alert Message": alert.alert_message,
                "Suggested Action": alert.suggested_action or "N/A",
                "Priority": alert.priority,
                "Read Status": "Read" if alert.is_read else "Unread",
                "Resolved": "Yes" if alert.is_resolved else "No",
                "Resolved Date": alert.resolved_at.strftime("%Y-%m-%d %H:%M") if alert.resolved_at else "N/A",
                "Created Date": alert.created_at.strftime("%Y-%m-%d %H:%M")
            })
        
        df = pd.DataFrame(data)
        
        if format_type == "csv":
            filename = f"alerts_export_{timestamp}.csv"
            filepath = os.path.join(self.export_dir, filename)
            df.to_csv(filepath, index=False)
        else:  # Excel
            filename = f"alerts_export_{timestamp}.xlsx"
            filepath = os.path.join(self.export_dir, filename)
            
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Alerts', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Alerts']
                
                # Format header
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid")
                
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add summary sheet
                critical_alerts = len([a for a in alerts if a.priority == "critical"])
                high_alerts = len([a for a in alerts if a.priority == "high"])
                unread_alerts = len([a for a in alerts if not a.is_read])
                resolved_alerts = len([a for a in alerts if a.is_resolved])
                
                summary_data = {
                    "Metric": [
                        "Total Alerts",
                        "Critical Alerts",
                        "High Priority Alerts",
                        "Medium Priority Alerts",
                        "Low Priority Alerts",
                        "Unread Alerts",
                        "Resolved Alerts"
                    ],
                    "Value": [
                        len(alerts),
                        critical_alerts,
                        high_alerts,
                        len([a for a in alerts if a.priority == "medium"]),
                        len([a for a in alerts if a.priority == "low"]),
                        unread_alerts,
                        resolved_alerts
                    ]
                }
                
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # Format summary sheet
                summary_sheet = writer.sheets['Summary']
                for cell in summary_sheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
        
        return filepath